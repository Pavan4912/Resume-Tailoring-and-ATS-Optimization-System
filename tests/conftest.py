import os
from pathlib import Path
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def build_resume(path: Path, *, extra_bullet: str = "") -> Path:
    d = Document()
    d.add_paragraph("Asha Test")
    d.add_paragraph("Pune, India | +91 9000000000 | asha@example.com")
    d.add_paragraph("PROFESSIONAL SUMMARY", style="Heading 2")
    p = d.add_paragraph()
    for chunk in ["Java developer with ", "2 years", " of experience building ", "REST APIs", " using Spring Boot."]:
        p.add_run(chunk)                                  # fragmented runs
    d.add_paragraph("TECHNICAL SKILLS", style="Heading 2")
    p = d.add_paragraph()
    p.add_run("Backend:").bold = True
    p.add_run(" Java, ")
    p.add_run("Spring Boot")
    p.add_run(", Hibernate")
    p = d.add_paragraph()                                  # two lines joined by a soft break
    p.add_run("Tools:").bold = True
    p.add_run(" Maven, Git")
    p.add_run().add_break()
    p.add_run("Frontend: ").bold = True      # trailing space inside the bold run, as in real resumes
    p.add_run("HTML, React")
    d.add_paragraph("PROFESSIONAL EXPERIENCE", style="Heading 2")
    p = d.add_paragraph()
    p.add_run("Software Engineer | Acme Corp \u2014 Pune\t")
    p.add_run("Jan 2022 \u2013 Present").bold = True
    for t in ["Built REST APIs with Spring Boot and Hibernate.",
              "Improved database queries and fixed production issues."] + ([extra_bullet] if extra_bullet else []):
        d.add_paragraph(t, style="List Bullet")
    d.add_paragraph("EDUCATION", style="Heading 2")
    d.add_paragraph("B.E. Information Technology | Test University | Jun 2017 \u2013 May 2021")
    # hyperlink paragraph (cannot be edited safely)
    p = d.add_paragraph()
    p._p.append(parse_xml(f'<w:hyperlink {nsdecls("w")} w:anchor="x"><w:r><w:t>link text</w:t></w:r></w:hyperlink>'))
    d.save(path)
    return path


@pytest.fixture
def synth(tmp_path):
    return build_resume(tmp_path / "resume.docx")


@pytest.fixture
def real_master():
    p = os.environ.get("RESUME_TEST_MASTER")
    if not p or not Path(p).is_file():
        pytest.skip("set RESUME_TEST_MASTER to run tests against a real resume")
    return Path(p)
